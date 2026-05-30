from __future__ import annotations

import re
from io import BytesIO
from pathlib import Path

from issuesense.engiagent import build_investigation_draft


TEXT_EXTENSIONS = {".txt", ".md", ".log", ".csv", ".json"}
DOCUMENT_EXTENSIONS = TEXT_EXTENSIONS | {".pdf", ".docx"}

SIGNAL_TERMS = [
    "error",
    "exception",
    "failed",
    "failure",
    "timeout",
    "latency",
    "expected",
    "actual",
    "staging",
    "production",
    "regression",
    "requirement",
    "api",
    "database",
    "memory",
    "cpu",
]

INCIDENT_TERMS = [
    "actual:",
    "expected:",
    "failed on",
    "fails on",
    "failure is reproducible",
    "traceback",
    "stack trace",
    "http 500",
    "server error",
    "crash",
    "cors policy",
    "failed to fetch",
    "blocked by cors",
]

ENGINEERING_CONTEXT_TERMS = [
    "api",
    "backend",
    "frontend",
    "database",
    "requirement",
    "acceptance criteria",
    "test report",
    "test case",
    "incident",
    "root cause",
    "8d",
    "traceback",
    "stack trace",
    "deployment",
    "regression",
    "staging",
    "production",
    "service",
    "connector",
    "workflow",
    "validation",
    "qa",
    "defect",
    "latency",
    "timeout",
    "http",
    "error",
]

NON_ENGINEERING_TERMS = [
    "unikey",
    "tone0",
    "tone1",
    "tone2",
    "tone3",
    "tone4",
    "tone5",
    "key mapping",
    "hook-bowl",
]


def analyze_engineering_document(filename: str, content: bytes, triage: dict | None = None) -> dict:
    text = extract_document_text(filename, content)
    chunks = chunk_document_text(text)
    summary = summarize_document(filename, text, chunks)
    evidence = extract_document_evidence(chunks)
    relevance = assess_document_relevance(filename, text, summary)
    if not relevance["accepted"]:
        return {
            "module": "EngiAgent",
            "document": build_document_metadata(filename, content, text, chunks),
            "summary": summary,
            "document_relevance": relevance,
            "document_triage": {
                "predicted_label": "unsupported_document",
                "likely_cause": relevance["reason"],
                "source": "document_relevance_guardrail",
                "actionable_incident": False,
            },
            "chunks": chunks,
            "evidence": evidence[:2],
            "investigation": build_unsupported_document_response(summary, relevance),
        }
    resolved_triage = normalize_document_triage(triage, summary, text)
    investigation_text = build_investigation_input(summary, evidence, text)
    investigation = build_investigation_draft(investigation_text, resolved_triage)

    return {
        "module": "EngiAgent",
        "document": build_document_metadata(filename, content, text, chunks),
        "summary": summary,
        "document_relevance": relevance,
        "document_triage": resolved_triage,
        "chunks": chunks,
        "evidence": evidence,
        "investigation": investigation,
    }


def build_document_metadata(filename: str, content: bytes, text: str, chunks: list[dict]) -> dict:
    return {
        "filename": filename,
        "extension": Path(filename).suffix.lower(),
        "bytes": len(content),
        "characters": len(text),
        "chunk_count": len(chunks),
    }


def extract_document_text(filename: str, content: bytes) -> str:
    extension = Path(filename).suffix.lower()
    if extension not in DOCUMENT_EXTENSIONS:
        raise ValueError(
            f"Unsupported document type '{extension or 'unknown'}'. "
            "Use .txt, .md, .log, .csv, .json, .pdf, or .docx."
        )
    if not content:
        raise ValueError("Uploaded document is empty.")
    if extension in TEXT_EXTENSIONS:
        return _decode_text(content)
    if extension == ".pdf":
        return _extract_pdf_text(content)
    if extension == ".docx":
        return _extract_docx_text(content)
    raise ValueError(f"Unsupported document type '{extension}'.")


def chunk_document_text(text: str, chunk_size: int = 950) -> list[dict]:
    clean_text = normalize_document_text(text)
    if not clean_text:
        raise ValueError("No readable text was extracted from the document.")

    paragraphs = [part.strip() for part in re.split(r"\n\s*\n", clean_text) if part.strip()]
    chunks: list[dict] = []
    buffer = ""
    start = 0
    cursor = 0

    for paragraph in paragraphs:
        if not buffer:
            start = cursor
        candidate = f"{buffer}\n\n{paragraph}".strip() if buffer else paragraph
        if len(candidate) > chunk_size and buffer:
            chunks.append(_make_chunk(len(chunks) + 1, buffer, start))
            start = cursor
            buffer = paragraph
        else:
            buffer = candidate
        cursor += len(paragraph) + 2

    if buffer:
        chunks.append(_make_chunk(len(chunks) + 1, buffer, start))

    return chunks


def summarize_document(filename: str, text: str, chunks: list[dict]) -> dict:
    lines = [line.strip() for line in normalize_document_text(text).splitlines() if line.strip()]
    signal_lines = [line for line in lines if _signal_score(line) > 0]
    key_findings = [_compact(line, 220) for line in signal_lines[:6]]
    if not key_findings:
        key_findings = [_compact(line, 220) for line in lines[:4]]

    return {
        "title": Path(filename).stem or filename,
        "document_type": infer_document_type(filename, text),
        "word_count": len(re.findall(r"\b\w+\b", text)),
        "chunk_count": len(chunks),
        "key_findings": key_findings,
        "risk_level": infer_risk_level(text),
    }


def extract_document_evidence(chunks: list[dict], limit: int = 6) -> list[dict]:
    ranked = sorted(chunks, key=lambda chunk: _signal_score(chunk["text"]), reverse=True)
    evidence = []
    for chunk in ranked[:limit]:
        score = _signal_score(chunk["text"])
        evidence.append(
            {
                "id": f"DOC-EV-{len(evidence) + 1:02d}",
                "chunk_id": chunk["id"],
                "score": score,
                "text": _compact(chunk["text"], 360),
                "source": "uploaded_document",
            }
        )
    return evidence


def build_investigation_input(summary: dict, evidence: list[dict], text: str) -> str:
    findings = "\n".join(f"- {item}" for item in summary["key_findings"])
    evidence_text = "\n".join(f"- {item['text']}" for item in evidence[:4])
    if not findings and not evidence_text:
        return _compact(text, 1800)
    return (
        f"Document type: {summary['document_type']}.\n"
        f"Risk level: {summary['risk_level']}.\n"
        f"Key findings:\n{findings}\n"
        f"Evidence snippets:\n{evidence_text}"
    )


def infer_document_type(filename: str, text: str) -> str:
    lowered = f"{filename} {text}".lower()
    name = Path(filename).name.lower()
    if name in {"readme.md", "readme.txt"} or any(
        term in lowered for term in ["quick start", "portfolio claims", "available endpoints", "project structure"]
    ):
        return "project_document"
    if any(term in lowered for term in ["traceback", "exception", "console", "stack", "cors policy", "failed to fetch"]):
        return "runtime_log"
    if any(term in lowered for term in ["test report", "test execution", "actual:", "expected:"]):
        return "test_report"
    if any(term in lowered for term in ["incident", "root cause", "8d", "containment"]):
        return "incident_report"
    if any(term in lowered for term in ["requirement", "acceptance criteria", "user story"]):
        return "requirement_document"
    return "engineering_note"


def infer_risk_level(text: str) -> str:
    lowered = text.lower()
    high_terms = ["production outage", "data loss", "security breach", "critical failure", "cannot proceed"]
    medium_terms = ["staging fails", "regression failure", "timeout", "failed", "http 500", "blocked"]
    if any(term in lowered for term in high_terms):
        return "high"
    if any(term in lowered for term in medium_terms):
        return "medium"
    return "low"


def infer_document_triage(summary: dict, text: str) -> dict:
    lowered = text.lower()
    incident_score = sum(1 for term in INCIDENT_TERMS if term in lowered)
    if summary["document_type"] in {"test_report", "incident_report", "runtime_log"} and incident_score >= 1:
        return {
            "predicted_label": _infer_incident_label(lowered),
            "likely_cause": "document contains failure or validation evidence that should be investigated",
            "source": "document_intake_inference",
            "actionable_incident": True,
        }

    return {
        "predicted_label": "document_review",
        "likely_cause": "uploaded document is reference or requirement context rather than a confirmed failure",
        "source": "document_intake_inference",
        "actionable_incident": False,
    }


def assess_document_relevance(filename: str, text: str, summary: dict) -> dict:
    lowered = f"{filename} {text}".lower()
    non_engineering_score = sum(1 for term in NON_ENGINEERING_TERMS if term in lowered)
    engineering_score = sum(1 for term in ENGINEERING_CONTEXT_TERMS if term in lowered)
    accepted_types = {
        "test_report",
        "incident_report",
        "requirement_document",
        "runtime_log",
        "project_document",
    }

    if non_engineering_score >= 2 and engineering_score < 3:
        return {
            "accepted": False,
            "score": engineering_score,
            "reason": "document looks like a keyboard/config mapping, not an engineering issue, requirement, log, or test report",
            "expected_inputs": expected_document_inputs(),
        }
    if summary["document_type"] in accepted_types:
        return {
            "accepted": True,
            "score": engineering_score,
            "reason": f"accepted as {summary['document_type']}",
            "expected_inputs": expected_document_inputs(),
        }
    if engineering_score >= 3:
        return {
            "accepted": True,
            "score": engineering_score,
            "reason": "accepted because engineering workflow terms were detected",
            "expected_inputs": expected_document_inputs(),
        }
    return {
        "accepted": False,
        "score": engineering_score,
        "reason": "document does not contain enough engineering workflow context for EngiAgent",
        "expected_inputs": expected_document_inputs(),
    }


def expected_document_inputs() -> list[str]:
    return [
        "test reports with expected/actual results",
        "incident notes with observed failure and reproduction context",
        "runtime logs or stack traces",
        "requirements or acceptance criteria",
        "engineering design notes with workflow/API/system constraints",
    ]


def build_unsupported_document_response(summary: dict, relevance: dict) -> dict:
    return {
        "module": "EngiAgent",
        "status": "document_rejected",
        "agent_runtime": "document_intake_guardrail",
        "predicted_label": "unsupported_document",
        "agent_plan": [
            "Validate document relevance.",
            "Skip investigation workflow when the input is outside the engineering domain.",
        ],
        "investigation_summary": (
            "EngiAgent skipped this upload because it does not look like an engineering issue, "
            "test report, runtime log, requirement, or design note."
        ),
        "review_artifact": {
            "review_type": "unsupported_document",
            "extracted_context": " ".join(summary["key_findings"][:3]),
            "suggested_use": "Upload an engineering artifact before running investigation or document review.",
            "open_questions": [
                "Is there a concrete issue, requirement, test result, or runtime failure to analyze?",
                "Can you provide expected behavior, actual behavior, reproduction steps, or acceptance criteria?",
            ],
            "next_steps": [
                "Use a test report, incident note, stack trace, requirement document, or engineering design note.",
                "Do not route keyboard maps, unrelated configs, or personal text files into the agent workflow.",
            ],
            "expected_inputs": relevance["expected_inputs"],
        },
        "eight_d": {},
        "evidence": [],
        "tool_trace": [
            {"tool": "document_relevance_guardrail", "result": relevance["reason"]},
            {"tool": "agent_workflow_router", "result": "EngiAgent workflow skipped"},
        ],
        "memory_notes": ["No investigation memory was created for unsupported documents."],
        "guardrails": {
            "accepted_engineering_document": False,
            "grounded_in_submitted_report": False,
            "complete_8d_fields": None,
            "missing_fields": [],
            "requires_human_approval": False,
        },
    }


def normalize_document_triage(triage: dict | None, summary: dict, text: str) -> dict:
    inferred = infer_document_triage(summary, text)
    if not triage:
        return inferred
    normalized = {**triage}
    normalized.setdefault(
        "predicted_label",
        normalized.get("label") or normalized.get("category") or inferred["predicted_label"],
    )
    normalized.setdefault("likely_cause", inferred["likely_cause"])
    normalized.setdefault("source", "provided_triage")
    normalized.setdefault("actionable_incident", inferred["actionable_incident"])
    return normalized


def normalize_document_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _extract_pdf_text(content: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise ValueError("PDF support requires pypdf. Install it with: pip install pypdf") from exc

    reader = PdfReader(BytesIO(content))
    pages = [page.extract_text() or "" for page in reader.pages]
    return normalize_document_text("\n\n".join(pages))


def _extract_docx_text(content: bytes) -> str:
    try:
        from docx import Document
    except ImportError as exc:
        raise ValueError("DOCX support requires python-docx. Install it with: pip install python-docx") from exc

    document = Document(BytesIO(content))
    lines = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                lines.append(" | ".join(cells))
    return normalize_document_text("\n".join(lines))


def _decode_text(content: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return normalize_document_text(content.decode(encoding))
        except UnicodeDecodeError:
            continue
    return normalize_document_text(content.decode("utf-8", errors="replace"))


def _make_chunk(index: int, text: str, start: int) -> dict:
    clean = text.strip()
    return {
        "id": f"CHUNK-{index:03d}",
        "text": clean,
        "start_char": start,
        "end_char": start + len(clean),
        "signal_score": _signal_score(clean),
    }


def _signal_score(text: str) -> int:
    lowered = text.lower()
    return sum(1 for term in SIGNAL_TERMS if term in lowered)


def _infer_incident_label(lowered: str) -> str:
    if any(term in lowered for term in ["cors policy", "failed to fetch", "blocked by cors", "network response"]):
        return "test_environment_issue"
    if any(term in lowered for term in ["latency", "timeout", "slow", "memory", "cpu"]):
        return "performance_issue"
    if any(term in lowered for term in ["staging", "environment", "config", "sandbox", "dependency"]):
        return "test_environment_issue"
    if any(term in lowered for term in ["expected", "acceptance criteria", "requirement", "undefined"]):
        return "requirement_gap"
    if any(term in lowered for term in ["api contract", "webhook", "connector", "sync", "field name"]):
        return "integration_issue"
    if any(term in lowered for term in ["duplicate", "missing data", "null", "csv", "schema"]):
        return "data_issue"
    return "software_bug"


def _compact(text: str, limit: int) -> str:
    clean = " ".join(text.split())
    if len(clean) <= limit:
        return clean
    return f"{clean[: limit - 3]}..."
